from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "StreamForge_Project_Explained.docx"
NAVY = "17243A"
PALE_BLUE = "EAF1F8"
PALE_GREEN = "E8F4ED"
PALE_GOLD = "FBF2DE"
MID_GREY = "596579"
GRID = "D7DFE9"


def shade(cell, fill):
    props = cell._tc.get_or_add_tcPr()
    shd = props.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        props.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=GRID, size="5"):
    props = cell._tc.get_or_add_tcPr()
    borders = props.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        props.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_text(cell, text, *, bold=False, color=NAVY, size=9.3):
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)


def format_table(table, widths=None, header=True):
    table.autofit = False
    for row_idx, row in enumerate(table.rows):
        tr_pr = row._tr.get_or_add_trPr()
        tr_pr.append(OxmlElement("w:cantSplit"))
        for col_idx, cell in enumerate(row.cells):
            if widths:
                cell.width = Inches(widths[col_idx])
            set_cell_border(cell)
            if header and row_idx == 0:
                shade(cell, NAVY)
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.color.rgb = RGBColor(255, 255, 255)
                        run.bold = True
            elif row_idx % 2 == 0:
                shade(cell, "F5F8FB")


def add_table(document, headers, rows, widths=None):
    table = document.add_table(rows=1, cols=len(headers))
    for cell, label in zip(table.rows[0].cells, headers):
        set_cell_text(cell, label, bold=True, color="FFFFFF", size=9)
    for values in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            set_cell_text(cell, str(value))
    format_table(table, widths=widths)
    return table


def add_callout(document, title, body, fill=PALE_BLUE):
    table = document.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    shade(cell, fill)
    set_cell_border(cell, color=fill, size="0")
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title)
    r.bold = True
    r.font.name = "Aptos"
    r.font.size = Pt(10.5)
    r.font.color.rgb = RGBColor.from_string(NAVY)
    p2 = cell.add_paragraph(body)
    p2.paragraph_format.space_after = Pt(5)
    p2.paragraph_format.line_spacing = 1.05
    for run in p2.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor.from_string(NAVY)
    document.add_paragraph().paragraph_format.space_after = Pt(1)


def add_bullet(document, text):
    p = document.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.first_line_indent = Inches(-0.16)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.05
    p.add_run(text)
    return p


def add_code(document, text):
    p = document.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.right_indent = Inches(0.12)
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F2F4F7")
    pPr.append(shd)
    r = p.add_run(text)
    r.font.name = "Consolas"
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor.from_string(NAVY)


def add_heading(document, text, level=1):
    p = document.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    return p


def build():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.62)
    sec.bottom_margin = Inches(0.62)
    sec.left_margin = Inches(0.72)
    sec.right_margin = Inches(0.72)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor.from_string(NAVY)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08
    for name, size in (("Title", 27), ("Heading 1", 17), ("Heading 2", 12)):
        style = doc.styles[name]
        style.font.name = "Aptos Display" if name == "Title" else "Aptos"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(5 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(6)

    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.paragraph_format.space_before = Pt(2)
    fr = footer.add_run(f"StreamForge project guide  |  {date(2026, 10, 1).strftime('%B %Y')}")
    fr.font.name = "Aptos"
    fr.font.size = Pt(8)
    fr.font.color.rgb = RGBColor.from_string(MID_GREY)

    # Page 1: explain the purpose and give the current, observed status.
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(2)
    p.add_run("STREAMFORGE  /  PROJECT GUIDE").bold = True
    p.runs[0].font.size = Pt(9)
    p.runs[0].font.color.rgb = RGBColor.from_string(MID_GREY)
    title = doc.add_paragraph(style="Title")
    title.add_run("StreamForge Project Explained")
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(12)
    sr = subtitle.add_run("A plain-language guide to what the system does, what is working, and what comes next")
    sr.italic = True
    sr.font.size = Pt(11)
    sr.font.color.rgb = RGBColor.from_string(MID_GREY)

    add_callout(
        doc,
        "In one sentence",
        "StreamForge reads truck temperature messages from one or more Kafka topics, calculates a separate five-minute average for each truck in each stream, saves the results, and shows live system health in a web dashboard.",
        PALE_GREEN,
    )
    add_heading(doc, "What is working now", 1)
    doc.add_paragraph(
        "The web dashboard is connected to the running Kafka cluster and the StreamForge API. "
        "The current local setup has 20 processor workers. Kafka shares 40 input partitions from two topics between them, "
        "and all workers report healthy local RocksDB state stores. Each worker currently owns two partitions, one from each topic."
    )
    add_table(doc, ["Live check", "Observed result"], [
        ("Kafka", "Connected; 1 broker"),
        ("Input topics", "2 independent streams; 20 partitions each"),
        ("Processor workers", "20 active; 2 partitions per worker across the two topics"),
        ("Consumer lag", "0 events at the time checked"),
        ("RocksDB", "20 of 20 worker stores healthy"),
        ("Resource monitoring", "CPU and memory values are being reported per worker"),
    ], widths=[2.05, 4.95])
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    rr = p.add_run("The live values above are a snapshot, not fixed project targets. They can change as events arrive and workers run.")
    rr.italic = True
    rr.font.size = Pt(9)
    rr.font.color.rgb = RGBColor.from_string(MID_GREY)

    # Page 2: give a straightforward pipeline explanation.
    doc.add_page_break()
    add_heading(doc, "What the system does", 1)
    doc.add_paragraph(
        "A truck sends a temperature reading with a truck ID, event ID, temperature, and timestamp. "
        "Kafka holds the messages in one or more input topics, which are independent event streams. The processor reads all configured topics in batches of up to 128 messages, checks required fields, "
        "and groups valid readings into five-minute tumbling windows for each truck within its source stream."
    )
    add_table(doc, ["Step", "What happens"], [
        ("1. Receive", "A producer writes each truck reading to one of the configured Kafka input topics."),
        ("2. Check", "The processor checks the message fields and timestamp. Readings at or below 0°C are filtered out."),
        ("3. Calculate", "For each input topic, truck, and five-minute tumbling window, it updates the running total, count, and average."),
        ("4. Save", "The new average goes to truck-temperature-averages with its source topic. A matching namespaced state snapshot goes to a compacted changelog topic."),
        ("5. Recover", "Each worker keeps its own RocksDB files and can rebuild state from committed changelog snapshots."),
        ("6. Display", "The API reads Kafka, worker, and metrics information. The React dashboard refreshes it about every five seconds."),
    ], widths=[1.25, 5.75])
    add_callout(
        doc,
        "Why use transactions",
        "The output average, saved state snapshot, and Kafka input offsets for a batch are committed together. RocksDB is updated after the Kafka transaction commits. This helps prevent an input batch from being marked finished when its results were not saved.",
    )
    add_heading(doc, "The main parts", 2)
    add_table(doc, ["Part", "Plain-language role"], [
        ("Kafka", "Moves and stores the incoming and outgoing event streams."),
        ("Processor", "Validates readings, calculates averages, and writes results."),
        ("RocksDB", "Stores each worker's local copy of the calculated window state."),
        ("FastAPI", "Provides live cluster and worker information to the dashboard."),
        ("React dashboard", "Shows cluster status, partitions, workers, lag, resource use, and topology."),
    ], widths=[1.75, 5.25])
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    r = p.add_run("Implementation note: the processor currently uses the Confluent Kafka Python client directly; it is not running the Faust framework.")
    r.italic = True
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor.from_string(MID_GREY)

    # Page 3: summarize implementation and practical outcomes.
    doc.add_page_break()
    add_heading(doc, "What has been built", 1)
    add_table(doc, ["Area", "Completed work"], [
        ("Real data in the UI", "Removed mock dashboard data and connected pages to the live API."),
        ("Kafka status", "The API reads broker and topic metadata, partition leaders, assignments, and consumer lag."),
        ("Multiple workers", "Workers use unique Kafka transaction IDs and separate RocksDB directories. Kafka assigns partitions within one shared consumer group."),
        ("Multiple input streams", "The processor can subscribe to a comma-separated list of Kafka topics. Output records and state keys keep the source topic, so matching truck IDs in different streams stay separate."),
        ("State recovery", "Workers restore committed snapshots at startup and before taking newly assigned input partitions."),
        ("RocksDB health", "Each worker exports whether its RocksDB store is open; the dashboard shows how many worker stores are healthy."),
        ("CPU and memory", "Each worker samples its own process CPU and resident memory about every two seconds and publishes the readings as metrics."),
        ("Build and setup", "The repository includes local setup steps, Kafka topic creation, API and UI startup commands, and project notes."),
    ], widths=[1.55, 5.45])
    add_heading(doc, "How to read resource use", 2)
    doc.add_paragraph(
        "CPU is shown as the worker's share of all logical CPU cores on the host computer. "
        "Memory is the worker process's resident working set as a share of the computer's physical memory. "
        "A low or zero CPU reading is normal when a worker has no messages to process. These are live measurements, not reserved capacity."
    )
    add_heading(doc, "How worker count relates to partitions", 2)
    doc.add_paragraph(
        "A partition is a piece of a Kafka topic; it is not a worker by itself. The live setup has 40 partitions and 20 workers, "
        "so Kafka currently assigns two partitions to each worker. Adding more workers does not create more partitions, and workers beyond the partition count would have no input partition to process."
    )
    add_callout(
        doc,
        "Current live check",
        "At the latest live check, the API returned 20 workers, 40 assigned partitions, zero consumer lag, and healthy RocksDB stores on all workers.",
        PALE_GREEN,
    )

    # Page 4: concise run instructions and verification evidence.
    doc.add_page_break()
    add_heading(doc, "How to start the local project", 1)
    doc.add_paragraph("Run these commands from the project folder in PowerShell. Docker Desktop must be running for the Kafka containers.")
    add_heading(doc, "Start Kafka and prepare topics", 2)
    add_code(doc, "docker compose -f kafka/docker-compose.yml up -d")
    add_code(doc, ".venv\\Scripts\\python.exe kafka/topic_setup.py")
    add_heading(doc, "Start 20 processor workers", 2)
    doc.add_paragraph("Prepare both input topics. Start one processor process per metrics port, from 9101 through 9120:")
    add_code(doc, ".venv\\Scripts\\python.exe kafka/topic_setup.py --input-topics truck-telemetry,truck-telemetry-secondary")
    add_code(doc, ".venv\\Scripts\\python.exe -m faust.processor --input-topics truck-telemetry,truck-telemetry-secondary --batch-size 128 --metrics-port 9101")
    add_code(doc, ".venv\\Scripts\\python.exe -m faust.processor --input-topics truck-telemetry,truck-telemetry-secondary --batch-size 128 --metrics-port 9102  # repeat through port 9120")
    add_heading(doc, "Start the API and dashboard", 2)
    add_code(doc, '$env:KAFKA_INPUT_TOPICS = "truck-telemetry,truck-telemetry-secondary"')
    add_code(doc, ".venv\\Scripts\\python.exe -m uvicorn backend.api:app --reload --port 8000")
    add_code(doc, "cd ui")
    add_code(doc, "npm start")
    doc.add_paragraph("Open http://localhost:3000 in a browser. The dashboard polls the API every five seconds.")
    add_heading(doc, "Verification completed", 2)
    add_bullet(doc, "Python source files compiled successfully after the latest changes.")
    add_bullet(doc, "The React production build completed successfully.")
    add_bullet(doc, "The live API reported Kafka connected, 20 processor workers, 40 total partitions across two topics, and zero lag.")
    add_bullet(doc, "All 20 worker metric endpoints returned CPU, memory, per-worker event counters, processor, and RocksDB health metrics.")
    add_bullet(doc, "The dashboard data showed 20 healthy RocksDB stores and live resource values for each worker.")
    add_bullet(doc, "The processor and dashboard now support multiple input topics; per-topic state and output fields keep the streams distinct.")
    add_bullet(doc, "A same-truck smoke check produced separate 11°C and 31°C averages from the two input topics, confirming that their state was not combined.")
    add_bullet(doc, "A 100,000-message two-topic run delivered all messages without producer failures. The processor reached zero lag; measured valid-event processing is below the 100,000-events-per-second target.")
    add_bullet(doc, "Measured load run: 100,000 input messages were accounted for as 83,362 processed, 16,638 filtered, and 0 errors. Kafka delivery and processing drained in 11.52 seconds, about 8,681 input messages per second end to end, with zero lag.")
    add_bullet(doc, "During a second run of 18,001 messages, one worker was terminated. The Kafka group rebalanced; after the worker restarted, the cluster returned to 20 workers, zero lag, and 20 healthy state stores. The restarted worker logged restoration of 181,853 changelog snapshots.")
    add_callout(
        doc,
        "What provides the sample traffic",
        "The repository includes a producer that can send example truck readings. A real truck device or external event source still needs to publish messages using the project's JSON fields.",
        PALE_GOLD,
    )

    # Page 5: clearly separate known limits from practical next steps.
    doc.add_page_break()
    add_heading(doc, "Limits and next steps", 1)
    add_heading(doc, "Known limits", 2)
    add_bullet(doc, "The local Docker setup uses one Kafka broker with replication factor one. It is suitable for a local project demo but does not provide broker redundancy.")
    add_bullet(doc, "The specified 100,000 events-per-second processing target has not been demonstrated. The measured end-to-end run handled about 8,681 input messages per second on this local setup; the standalone producer reached about 33,614 messages per second.")
    add_bullet(doc, "The worker termination and restart check verified rebalance, changelog restoration, and return to zero lag. It was a local demonstration, not exhaustive crash or network-failure testing. The Kafka setup has one broker with replication factor one.")
    add_bullet(doc, "The five-minute calculation uses tumbling windows. A continuously sliding rolling average and late-arriving-event policy are not implemented.")
    add_bullet(doc, "The project uses the Confluent Kafka Python client directly; it does not use Faust or Bytewax, as listed in the project brief.")
    add_bullet(doc, "The changelog currently holds more than 181,000 state snapshots in this local run. State/window cleanup and a retention policy are needed to bound disk use as the system runs longer.")
    add_bullet(doc, "Invalid messages are counted and skipped. A dead-letter topic for inspecting rejected messages is not yet included.")
    add_bullet(doc, "Worker CPU and memory are measured on the same host where the processor runs. Remote workers need their metrics endpoints listed in STREAMFORGE_PROCESSOR_METRICS_URLS.")
    add_bullet(doc, "The sample producer creates demonstration data. A live application still needs an actual telemetry source and an agreed message contract.")
    add_heading(doc, "Recommended next steps", 2)
    add_table(doc, ["Next step", "Why it helps"], [
        ("Connect a real telemetry source", "Proves the full project using data from its intended producer."),
        ("Add a dead-letter topic", "Keeps malformed messages available for review instead of only counting and skipping them."),
        ("Repeat worker failure and restart checks", "Builds confidence in partition reassignment and recovery across different failure points."),
        ("Optimize and benchmark processing throughput", "Measures sustained processing rate against the 100,000 events-per-second target."),
        ("Choose Faust/Bytewax or document the client-based design", "Aligns the implementation with the framework choices in the project brief."),
        ("Add state expiration and changelog cleanup", "Bounds RocksDB and Kafka storage as old windows become unnecessary."),
        ("Add broker redundancy for deployment", "Reduces the risk that one broker failure stops the whole pipeline."),
    ], widths=[2.5, 4.5])
    add_heading(doc, "Short glossary", 2)
    add_table(doc, ["Word", "Meaning"], [
        ("Event", "One truck reading, including its temperature and timestamp."),
        ("Consumer group", "The set of processor workers that share topic partitions."),
        ("Lag", "How many input messages the processor group is behind."),
        ("Changelog", "A Kafka topic containing saved state snapshots used for recovery."),
        ("RocksDB", "The local database each worker uses to store window state."),
    ], widths=[1.6, 5.4])

    doc.core_properties.title = "StreamForge Project Explained"
    doc.core_properties.subject = "Plain-language explanation of the StreamForge event processing project"
    doc.core_properties.author = "StreamForge Project Team"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
