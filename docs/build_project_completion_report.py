from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(__file__).parent / "StreamForge_Project_Completion_Report.docx"
NAVY = "17365D"
BLUE = "2E75B6"
PALE_BLUE = "EAF2F8"
PALE_GRAY = "F3F5F7"
MID = "596675"
WHITE = "FFFFFF"


def set_cell_fill(cell, color):
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color)
    cell._tc.get_or_add_tcPr().append(shading)


def set_repeat_table_header(row):
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    row._tr.get_or_add_trPr().append(header)


def set_font(run, size=10, bold=False, color="202833", italic=False):
    run.font.name = "Arial"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def add_text(paragraph, text, **kwargs):
    return set_font(paragraph.add_run(text), **kwargs)


def set_cell_text(cell, text, *, bold=False, color="202833", size=9):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    add_text(p, text, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for index, title in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], title, bold=True, color=WHITE, size=9)
        set_cell_fill(table.rows[0].cells[index], NAVY)
    set_repeat_table_header(table.rows[0])
    for row_index, data in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(data):
            set_cell_text(cells[index], str(value), size=8.7)
            if row_index % 2 == 1:
                set_cell_fill(cells[index], PALE_GRAY)
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Inches(width)
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(9 if level == 1 else 6)
    p.paragraph_format.space_after = Pt(4)
    add_text(p, text, size=15 if level == 1 else 11.5, bold=True, color=NAVY if level == 1 else BLUE)
    return p


def add_para(doc, text, *, lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.08
    if lead and text.startswith(lead):
        add_text(p, lead, size=9.5, bold=True)
        add_text(p, text[len(lead):], size=9.5)
    else:
        add_text(p, text, size=9.5)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.first_line_indent = Inches(-0.14)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.05
    add_text(p, text, size=9.2)
    return p


def add_code(doc, text):
    for line in text.strip().splitlines():
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.15)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        add_text(p, line, size=8.2, color=NAVY)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_text(paragraph, "Page ", size=8, color=MID)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.62)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    section.header_distance = Inches(0.28)
    section.footer_distance = Inches(0.3)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9.5)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    for style_name in ("Heading 1", "Heading 2"):
        doc.styles[style_name].font.name = "Arial"
        doc.styles[style_name]._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        doc.styles[style_name]._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")

    header = section.header.paragraphs[0]
    add_text(header, "STREAM FORGE   /   PROJECT COMPLETION", size=8, bold=True, color=NAVY)
    footer = section.footer.paragraphs[0]
    add_text(footer, "Project report  •  01 October 2026", size=8, color=MID)
    add_page_number(footer)

    title = doc.add_paragraph(style="Title")
    title.paragraph_format.space_before = Pt(10)
    title.paragraph_format.space_after = Pt(3)
    add_text(title, "Stream Forge Project Completion Report", size=25, bold=True, color=NAVY)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(9)
    add_text(subtitle, "Weeks 1–4 implementation and validation", size=14, color=BLUE)

    add_table(doc, ["Project details", "Value"], [
        ("Project", "Stream Forge — Distributed Python Event Processor"),
        ("Reporting period", "Weeks 1–4; completion review dated 1 October 2026"),
        ("Prepared by", "Member A — Sidhanshu branch"),
        ("Audience", "Project evaluator / instructor"),
        ("Status", "Local end-to-end pipeline implemented and verified; production hardening remains"),
    ], [1.55, 5.5])

    add_heading(doc, "Executive summary")
    add_para(doc, "Stream Forge now has a runnable local pipeline from Kafka telemetry input through Python processing and RocksDB state to a React monitoring dashboard. The processor filters temperatures at or below 0°C and maintains UTC five-minute tumbling-window averages per truck. Kafka transactions commit each valid event’s output, changelog snapshot, and source offset together.")
    add_para(doc, "The implemented stack was exercised against a live single-broker Kafka cluster. In a bounded run, 100 generated telemetry records were consumed: 85 positive-temperature records produced aggregate updates and 15 records were filtered. The API reported 20 active partitions, one processor consumer, zero lag, and live Prometheus metrics. A processor restart replayed 94 changelog snapshots. The current scope is suitable for local demonstration and further team development; multi-process recovery and crash-during-transaction behavior still need a dedicated fault-injection run.")

    add_heading(doc, "Delivered outcomes")
    add_table(doc, ["Area", "Delivered", "Evidence"], [
        ("Kafka foundation", "Three 20-partition topics; compacted state changelog", "Live broker metadata and topic setup verified"),
        ("Telemetry producer", "Validated JSON events, unique IDs, idempotent Kafka delivery", "100/100 sample records delivered"),
        ("Stream processor", "Filter, window aggregate, transactional output and offsets", "85 aggregate outputs; 15 filtered"),
        ("State recovery", "RocksDB store rebuilt from Kafka changelog", "94 changelog snapshots replayed on restart"),
        ("Monitoring", "FastAPI cluster data and Prometheus exposition proxy", "20/20 active partitions; zero consumer lag"),
        ("Dashboard", "Live API data on overview, topology, workers, partitions and metrics", "React production build compiled successfully"),
        ("Automated checks", "Producer, decoding, window and state tests", "9 tests passed"),
    ], [1.3, 3.3, 2.45])

    doc.add_page_break()
    add_heading(doc, "1  System design")
    add_para(doc, "The repository provides a three-topic Kafka pipeline, a transactional Python consumer/producer, local RocksDB state, a FastAPI monitoring layer, and a React Flow UI. The processing code uses the Confluent Kafka Python client directly. Faust and Bytewax were listed in the original plan but are not dependencies of the delivered processor.")
    add_table(doc, ["Pipeline stage", "Topic / service", "Responsibility"], [
        ("Input", "truck-telemetry", "Accept keyed telemetry records; truck_id keeps a truck’s records on one partition."),
        ("Process", "faust.processor", "Validate event, apply temperature filter, update five-minute window state."),
        ("Output", "truck-temperature-averages", "Publish the current average and count for each accepted event."),
        ("Recovery", "truck-temperature-averages-changelog", "Store the latest state snapshot per truck/window key; topic cleanup policy is compact."),
        ("State", "RocksDB", "Persist per-window sum, count, average, and event timestamp locally."),
        ("Monitoring", "FastAPI + Prometheus client", "Expose broker, consumer-group, lag, throughput, and processing-duration measurements."),
        ("Presentation", "React + React Flow", "Poll the API every five seconds and display current cluster state."),
    ], [1.0, 2.45, 3.6])
    add_heading(doc, "Telemetry contract", 2)
    add_para(doc, "Each input event has event_id, truck_id, temperature, and an ISO-8601 timestamp with a timezone. The producer generates sample records across 50,000 truck IDs and temperatures from −10°C to 50°C. The processor rejects malformed records, skips temperatures at or below 0°C, and groups accepted values by truck and UTC window start.")
    add_table(doc, ["Output field", "Meaning"], [
        ("event_id", "Source event identifier"),
        ("truck_id", "Truck key used to partition related events"),
        ("window_start / window_end", "UTC boundaries of a five-minute tumbling window"),
        ("average_temperature", "Cumulative mean for accepted readings in this window"),
        ("count", "Number of accepted readings included in the mean"),
        ("last_event_timestamp", "Timestamp of the source reading that updated the snapshot"),
    ], [2.1, 4.95])
    add_heading(doc, "Partitioning and scaling", 2)
    add_para(doc, "The three topics use 20 partitions so the processor group can distribute work across consumers. The live local verification used one consumer assigned all 20 partitions. Additional group members can share the partitions, but coordinated state recovery across concurrent members has not yet been tested under reassignment.")

    doc.add_page_break()
    add_heading(doc, "2  Processing and recovery")
    add_heading(doc, "Per-record transaction", 2)
    add_para(doc, "For each source message, the processor validates the payload and calculates the next aggregate from the current RocksDB value. It opens a Kafka transaction, writes the new changelog snapshot and output record when the event passes the filter, adds the source offset to the transaction, and commits. Only after commit does it update the local RocksDB value. Filtered or malformed records still have their source offset committed so they do not block later records; malformed payloads increment an error metric and are logged.")
    add_table(doc, ["Failure point", "Result"], [
        ("Before transaction commit", "Kafka hides the uncommitted output and changelog; the source offset is not committed and the record can be retried."),
        ("After transaction commit, before local RocksDB write", "Kafka output, changelog, and offset are committed together; startup replay restores the local state."),
        ("After local state update", "Kafka and RocksDB both contain the committed snapshot; normal processing continues from the next offset."),
    ], [2.2, 4.85])
    add_para(doc, "The exactly-once guarantee here applies to the Kafka transaction boundary: downstream consumers must use read_committed to hide aborted transactional records. RocksDB is not enlisted in the Kafka transaction; the compacted changelog is the recovery source for that separate local store.")
    add_heading(doc, "State model", 2)
    add_para(doc, "A RocksDB key is truck_id plus the UTC window-start epoch. The value stores the running sum and count, the derived average, the five-minute end time, and the latest event timestamp. The same JSON snapshot is published under that key to the compacted changelog. At process startup, the processor reads committed changelog records to the captured broker watermarks, applies snapshots, flushes RocksDB, and only then joins the input consumer group.")
    add_heading(doc, "Metrics", 2)
    add_table(doc, ["Metric", "Definition"], [
        ("streamforge_events_processed_total", "Accepted source events committed to the output path."),
        ("streamforge_events_filtered_total", "Valid events excluded because temperature is at or below 0°C."),
        ("streamforge_processing_errors_total", "Malformed input records skipped and logged."),
        ("streamforge_last_processing_duration_seconds", "Latest per-record processing duration, including Kafka commit and state flush."),
        ("streamforge_processor_up", "One while this processor instance is running."),
    ], [3.5, 3.55])
    add_para(doc, "The processor exports Prometheus text format on port 9101. FastAPI proxies the exposition at port 8000 `/metrics` and serves structured cluster data at `/api/cluster`.")

    doc.add_page_break()
    add_heading(doc, "3  Backend and dashboard")
    add_heading(doc, "Monitoring API", 2)
    add_table(doc, ["Endpoint", "Purpose"], [
        ("GET /api/health", "API process health check."),
        ("GET /api/cluster", "Kafka broker and partition metadata, active processor group members and assignments, source lag, throughput sample, processor status, processed count, and latest processing duration."),
        ("GET /metrics", "Proxy the processor’s Prometheus exposition; returns unavailable when the processor metrics server is stopped."),
    ], [1.8, 5.25])
    add_para(doc, "The API reads the configured Kafka topic and bootstrap servers, queries consumer group membership and committed offsets, and samples processor metrics from port 9101. It returns 503 with an actionable message when Kafka is unreachable. CORS permits the local React development server by default.")
    add_heading(doc, "Frontend integration", 2)
    add_para(doc, "A shared React context polls `/api/cluster` every five seconds and provides loading/error state across all pages. Overview cards use measured throughput, worker count, group lag, and processing duration. The topology and partitions views use broker leaders, replicas, in-sync replicas, group assignment, and per-partition lag. The metrics page lists these measurements and the Prometheus-backed processor values.")
    add_bullet(doc, "Removed generated frontend mock arrays and hard-coded metrics. The old unused metrics fetcher and its fabricated fallback values were removed with the mock data module.")
    add_bullet(doc, "The header, overview, sidebar, and topology indicator share the same API connection state; disconnected states no longer use green Connected or LIVE labels.")
    add_bullet(doc, "Worker membership, assigned partitions, and lag use Kafka consumer-group data. Per-worker throughput, CPU, and memory stay blank because the processor does not export those measurements.")
    add_table(doc, ["Dashboard page", "Live data displayed"], [
        ("Overview", "Kafka connection, sampled throughput, worker count, consumer lag, processing duration."),
        ("Topology", "Broker count, topic name, partition state, processor group state, architecture graph."),
        ("Workers", "Active consumer group members, assigned partitions, lag, status."),
        ("Partitions", "Leader, replicas, in-sync replicas, assigned processor member, lag."),
        ("Metrics", "Processed event count, processing rate, duration, lag, per-partition broker metadata."),
    ], [1.55, 5.5])
    add_heading(doc, "Operational topology", 2)
    add_para(doc, "The dashboard’s graph identifies Kafka, the processor consumer group, RocksDB, processor metrics, and the monitoring UI. The graph is an architectural view: RocksDB health itself is not currently probed, so the UI labels that measurement as unavailable instead of reporting a false healthy state.")

    doc.add_page_break()
    add_heading(doc, "4  Verification results")
    add_table(doc, ["Check", "Observed result", "Interpretation"], [
        ("Kafka cluster", "One broker; 20/20 telemetry partitions active", "Local broker and partition metadata available."),
        ("Topic setup", "Input, output, and changelog topics each have 20 partitions", "Changelog cleanup policy reported `compact`."),
        ("Producer run", "100 produced; 100 delivered; 0 failures", "Sample input reached Kafka successfully."),
        ("Processor run", "85 aggregate updates; 15 filtered; 0 processing errors", "Filter and output path handled all 100 records."),
        ("Consumer group", "One active member assigned 20 partitions; lag 0", "Processor caught up with the input topic."),
        ("Recovery restart", "94 changelog snapshots replayed", "RocksDB restored from broker state after process restart."),
        ("Unit tests", "9 passed", "Producer contract, event parsing, five-minute boundary, aggregate average, RocksDB persistence, and changelog restore."),
        ("Frontend build", "Optimized React production build compiled successfully", "All frontend imports and JSX compiled."),
    ], [1.4, 2.55, 3.1])
    add_heading(doc, "Local startup", 2)
    add_para(doc, "Run the following services in separate terminals from the repository root (except the UI command):")
    add_code(doc, ".venv\\Scripts\\python.exe -m pip install -r requirements.txt\ndocker compose -f kafka/docker-compose.yml up -d\n.venv\\Scripts\\python.exe kafka\\topic_setup.py\n.venv\\Scripts\\python.exe -m faust.processor\n.venv\\Scripts\\python.exe -m uvicorn backend.api:app --reload --port 8000\ncd ui && npm install && npm start")
    add_para(doc, "Produce sample input in an additional terminal with `.venv\\Scripts\\python.exe kafka\\producer.py --duration 60 --rate 1000`. The UI is available at `http://localhost:3000`; the API is at `http://localhost:8000`; processor metrics are at `http://localhost:8000/metrics`.")
    add_heading(doc, "Test and audit commands", 2)
    add_code(doc, ".venv\\Scripts\\python.exe -m pytest -q -p no:cacheprovider --basetemp .venv\\test-tmp tests\\test_kafka_producer.py tests\\test_processor.py tests\\test_state_store.py\n.venv\\Scripts\\python.exe kafka\\throughput_audit.py --topic truck-telemetry --events 100000\ncd ui && npm run build")

    doc.add_page_break()
    add_heading(doc, "5  Scope, limitations, and next steps")
    add_heading(doc, "Completed scope", 2)
    add_para(doc, "The local demonstration scope is implemented: Kafka setup and producer, schema checks, transactional filter-and-window processing, RocksDB state, a compacted changelog, API monitoring endpoints, Prometheus-format processor metrics, and a real-data dashboard. The project can be started locally using the commands in this report and the repository README.")
    add_heading(doc, "Known limitations", 2)
    add_bullet(doc, "Fault recovery was checked by restarting the processor and replaying 94 changelog snapshots. A deliberate kill during an in-flight Kafka transaction and automated assertion of no duplicates/loss were not run.")
    add_bullet(doc, "The local setup uses one Kafka broker and one processor member. Multi-instance reassignment, concurrent local state restoration, broker loss, and production durability have not been validated.")
    add_bullet(doc, "CPU and memory values are not collected. Processor throughput is reported as a short sampling-window rate; initial rate is blank until a second sample arrives.")
    add_bullet(doc, "Malformed telemetry is logged, counted, and skipped by committing its source offset. There is no dead-letter topic or replay workflow for invalid payloads.")
    add_bullet(doc, "The producer creates sample telemetry. A real fleet source must publish the documented event schema to Kafka. Security, authentication, TLS, multi-broker replication, retention, and deployment automation are not configured.")
    add_heading(doc, "Recommended next work", 2)
    add_table(doc, ["Priority", "Work item", "Acceptance evidence"], [
        ("1", "Add a dead-letter topic and inspect/replay controls for malformed records.", "Invalid event appears once in DLQ; valid input processing continues."),
        ("2", "Automate crash-in-transaction and consumer-rebalance tests with an isolated Kafka test stack.", "Input/output/changelog counts reconcile after repeated forced restarts."),
        ("3", "Validate state restore per assigned partition with multiple processor instances.", "Scale-up/scale-down tests preserve per-truck window totals without duplicates."),
        ("4", "Add host resource collection and stale-metric detection.", "CPU/memory and metric freshness displayed from real collectors."),
        ("5", "Prepare deployment settings for replicated Kafka, credentials, TLS, health checks, and persistent state volumes.", "Documented deployment rehearsal and recovery procedure."),
    ], [0.65, 3.4, 3.0])
    add_heading(doc, "References", 2)
    add_para(doc, "Project status source: supplied Weeks 1–2 completion report. Implementation evidence: repository code, local Kafka/API run, automated tests, and React production build described above.")
    add_para(doc, "Confluent Kafka Python documentation, transaction offsets and producer APIs: https://docs.confluent.io/platform/current/clients/confluent-kafka-python/html/index.html")
    add_para(doc, "RocksDict project documentation, Python RocksDB binding and persistent dictionary usage: https://github.com/rocksdict/RocksDict")

    props = doc.core_properties
    props.title = "Stream Forge Project Completion Report"
    props.subject = "Weeks 1–4 implementation and validation"
    props.author = "Stream Forge Project Team"
    props.keywords = "Stream processing, Kafka, RocksDB, Prometheus, FastAPI, React"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
