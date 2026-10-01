from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(__file__).parent / "StreamForge_Weeks_1_2_Internship_Report.docx"
NAVY = "17365D"
BLUE = "2E75B6"
LIGHT_BLUE = "D9EAF7"
LIGHT_GRAY = "F2F5F8"
MID_GRAY = "5B6573"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    table_properties = table._tbl.tblPr
    table_width = table_properties.first_child_found_in("w:tblW")
    table_width.set(qn("w:w"), str(sum(widths)))
    table_width.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for column, width in zip(grid.gridCol_lst, widths):
        column.set(qn("w:w"), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_font(run, size=10.5, bold=False, color="000000", italic=False):
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def add_text(paragraph, text, **kwargs):
    run = paragraph.add_run(text)
    set_font(run, **kwargs)
    return run


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph()
    paragraph.style = f"Heading {level}"
    add_text(paragraph, text, size=15 if level == 1 else 12, bold=True, color=NAVY if level == 1 else BLUE)
    return paragraph


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(3)
    add_text(paragraph, text, size=10.5)


def add_body(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(7)
    paragraph.paragraph_format.line_spacing = 1.15
    add_text(paragraph, text, size=10.5)
    return paragraph


def add_metadata(doc):
    table = doc.add_table(rows=4, cols=2)
    set_table_geometry(table, [2160, 7200])
    metadata = [
        ("Project", "StreamForge - Distributed Python Event Processor"),
        ("Reporting period", "Week 1 and Week 2"),
        ("Prepared by", "[Your Name], Intern"),
        ("Report purpose", "Progress update for internship project review"),
    ]
    for row, (label, value) in zip(table.rows, metadata):
        set_cell_shading(row.cells[0], LIGHT_BLUE)
        for cell in row.cells:
            cell.paragraphs[0].paragraph_format.space_after = Pt(2)
            cell.paragraphs[0].paragraph_format.space_before = Pt(2)
        add_text(row.cells[0].paragraphs[0], label, size=10, bold=True, color=NAVY)
        add_text(row.cells[1].paragraphs[0], value, size=10)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_status_table(doc):
    table = doc.add_table(rows=1, cols=4)
    set_table_geometry(table, [1660, 3500, 2500, 1700])
    headers = ["Week", "Objective", "Status", "Evidence"]
    for cell, text in zip(table.rows[0].cells, headers):
        set_cell_shading(cell, NAVY)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_text(p, text, size=9.5, bold=True, color=WHITE)
    rows = [
        ("Week 1", "Kafka foundation and dashboard scaffold", "Completed", "Kafka Compose stack, mock producer, React Flow dashboard"),
        ("Week 2", "Kafka integration and performance-readiness", "Completed", "20-partition topics, configurable producer, setup and audit tooling"),
    ]
    for index, values in enumerate(rows):
        cells = table.add_row().cells
        for cell, text in zip(cells, values):
            if index % 2 == 0:
                set_cell_shading(cell, LIGHT_GRAY)
            cell.paragraphs[0].paragraph_format.space_after = Pt(3)
            cell.paragraphs[0].paragraph_format.space_before = Pt(3)
            add_text(cell.paragraphs[0], text, size=9.5, bold=(text == "Completed"))


def build_report():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.15
    for name, size, color, before, after in [
        ("Heading 1", 15, NAVY, 14, 6),
        ("Heading 2", 12, BLUE, 10, 4),
    ]:
        style = styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_text(header, "STREAMFORGE | INTERNSHIP PROGRESS REPORT", size=8.5, bold=True, color=MID_GRAY)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(footer, "Weeks 1-2 | Internal Project Update", size=8.5, color=MID_GRAY)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_before = Pt(12)
    title.paragraph_format.space_after = Pt(3)
    add_text(title, "StreamForge", size=26, bold=True, color=NAVY)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(16)
    add_text(subtitle, "Internship Project Progress Report - Weeks 1 and 2", size=14, color=MID_GRAY)

    add_metadata(doc)

    add_heading(doc, "Executive Summary")
    add_body(
        doc,
        "During the first two weeks of the StreamForge internship project, the team established a working Kafka-based foundation for a Python distributed event-processing platform. My primary contribution focused on the integration layer: making the local Kafka environment reliable, defining the shared Kafka topic contract, improving the IoT telemetry producer, and preparing repeatable verification and throughput-audit tooling."
    )
    add_body(
        doc,
        "The current system can accept simulated truck telemetry through Apache Kafka and distribute it over 20 partitions. The next implementation stage is to connect the stream-processing application so it can consume these events, calculate five-minute temperature averages, and publish the results to the output topic."
    )

    add_heading(doc, "Milestone Status")
    add_status_table(doc)

    add_heading(doc, "Week 1: Foundation Setup")
    add_heading(doc, "Work Completed", level=2)
    for item in [
        "Reviewed and organized the StreamForge repository structure, architecture overview, dependencies, and local setup documentation.",
        "Configured a local Apache Kafka environment using Docker Compose, including Zookeeper, Kafka, Schema Registry, and topic initialization.",
        "Created a Python mock telemetry producer using confluent-kafka to generate truck ID, temperature, timestamp, and event ID data.",
        "Established truck_id as the Kafka message key so readings from the same truck are routed consistently to the same Kafka partition.",
        "Supported the initial React Flow topology dashboard scaffold that represents Kafka, the stream processor, state store, and Prometheus components.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "Week 2: Kafka Integration and Performance Readiness")
    add_heading(doc, "Work Completed", level=2)
    for item in [
        "Configured Kafka with separate internal and host-accessible listeners, allowing Docker services to communicate through kafka:29092 and local Python applications through localhost:9092.",
        "Defined and verified the shared 20-partition Kafka topic contract: truck-telemetry for incoming sensor readings and truck-temperature-averages for processed window results.",
        "Added a topic setup utility that creates missing topics, verifies their partition count, and safely expands local topics when they contain fewer than 20 partitions.",
        "Improved the telemetry producer with batching, idempotent delivery settings, optional event-count and rate controls, compression, and JSON performance output.",
        "Added a consumer throughput-audit utility for measuring Kafka read performance once the stream processor is connected.",
        "Added telemetry-schema unit tests and updated the project README with reproducible startup, topic setup, production, and audit commands.",
        "Committed and pushed the completed Week 2 integration work to the feature/telemetry-producer GitHub branch.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "Validation Evidence")
    evidence = doc.add_table(rows=1, cols=3)
    set_table_geometry(evidence, [3000, 4200, 2160])
    for cell, text in zip(evidence.rows[0].cells, ["Validation activity", "Result", "Outcome"]):
        set_cell_shading(cell, NAVY)
        add_text(cell.paragraphs[0], text, size=9.5, bold=True, color=WHITE)
    records = [
        ("Kafka topic setup", "truck-telemetry expanded to 20 partitions; output topic verified at 20 partitions.", "Pass"),
        ("Producer delivery test", "1,000 events produced and delivered; 0 failures; approximately 100 events/second at the configured rate.", "Pass"),
        ("Telemetry contract check", "Generated messages contain event_id, truck_id, temperature, and UTC timestamp fields.", "Pass"),
        ("Compose validation", "Docker Compose configuration validated before Kafka stack startup.", "Pass"),
    ]
    for index, values in enumerate(records):
        cells = evidence.add_row().cells
        for cell, text in zip(cells, values):
            if index % 2 == 0:
                set_cell_shading(cell, LIGHT_GRAY)
            cell.paragraphs[0].paragraph_format.space_after = Pt(3)
            cell.paragraphs[0].paragraph_format.space_before = Pt(3)
            add_text(cell.paragraphs[0], text, size=9.3, bold=(text == "Pass"))

    add_heading(doc, "Current Technical Status")
    add_body(doc, "The ingestion layer is operational. Kafka is available locally, the input and output topics have the required partitioning, and simulated telemetry can be delivered without observed message failures in the completed test. The system is ready for the stream-processing component to consume truck-telemetry and publish rolling aggregates to truck-temperature-averages.")

    add_heading(doc, "Planned Next Steps")
    for item in [
        "Integrate the Week 2 Python stream processor with the truck-telemetry Kafka topic.",
        "Implement event-time five-minute window aggregation per truck and publish results to truck-temperature-averages.",
        "Run end-to-end validation from telemetry producer through processor output.",
        "Use the throughput-audit tooling to measure consumer and processor throughput against the project target.",
        "Begin state persistence and recovery work for the next development phase.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "Risks and Dependencies")
    add_body(doc, "The current Kafka setup is a single-broker local development environment, so it validates the project workflow but does not yet demonstrate production-level replication or cross-machine fault tolerance. The planned end-to-end performance benchmark also depends on the stream-processing service being available. These limitations are expected at the present project stage and will be addressed in later milestones.")

    doc.core_properties.title = "StreamForge Internship Progress Report - Weeks 1 and 2"
    doc.core_properties.author = "[Your Name]"
    doc.core_properties.subject = "Two-week internship project progress report"
    OUT.parent.mkdir(exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build_report()
